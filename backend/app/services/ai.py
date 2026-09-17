"""Gemini integration: resume parsing, category mapping, AI-paste job extraction (Chapter 8, 27)."""
from __future__ import annotations

import json
import logging
from functools import lru_cache
from typing import Optional

from pydantic import BaseModel

from app.config import settings

logger = logging.getLogger("app.ai")

MAX_JOB_TEXT_CHARS = 20_000

# NOTE: none of the fields below may have default values (including Field(default=...) /
# default_factory). The google-genai SDK's response_schema converter hard-rejects any Pydantic
# field with a default ("Default value is not supported in the response schema for the Gemini
# API") - these models are only ever populated FROM a Gemini structured-output response (never
# manually constructed elsewhere with partial data), so requiring every field is safe; Gemini's
# controlled generation always emits the complete schema, using null for unknown values.


class ResumeData(BaseModel):
    is_resume: bool
    full_name: Optional[str]
    phone: Optional[str]
    email: Optional[str]
    highest_education: Optional[str]
    course: Optional[str]
    skills: list[str]
    experience_years: float
    city_or_district: Optional[str]
    summary: Optional[str]
    suggested_category_slugs: list[str]


class CategoryMatch(BaseModel):
    slug: Optional[str]


class JobDraft(BaseModel):
    title: Optional[str]
    company: Optional[str]
    category_slug: str
    qualification: Optional[str]
    district: Optional[str]
    location_text: Optional[str]
    job_type: str
    salary: Optional[str]
    apply_link: Optional[str]
    last_date: Optional[str]
    description: Optional[str]


class JobDraftList(BaseModel):
    jobs: list[JobDraft]


# ---------------------------------------------------------------------------
# Prompts (Chapter 27)
# ---------------------------------------------------------------------------

_RESUME_PROMPT = """You are a resume parser for an Indian job-alert service for students and freshers.
Extract the candidate details from the attached resume. Resumes may be in English, Marathi or
Hindi; always answer in English. Do not invent data: use null when a field is missing.

Rules:
- is_resume = false if the document is clearly not a resume/CV (random photo, ID card, marksheet only,
  notes, blank page).
- highest_education: the highest completed or ongoing qualification in short form
  (examples: "10th", "12th Science", "ITI Fitter", "Diploma Mechanical", "B.Com", "BCA", "BE Computer",
  "MBA Finance").
- course: main stream / specialisation if different from highest_education, else null.
- skills: at most 12 short skills (tools, software, trades, languages like "Typing 30 wpm"), no sentences.
- experience_years: total paid work experience in years (internships count as 0.5 max). Freshers = 0.
- city_or_district: current city/district only if written.
- summary: 1-2 simple English lines describing the candidate.
- suggested_category_slugs: up to 5 slugs, best match first, chosen ONLY from this list:
{category_list}

Return only JSON matching the schema."""

_CATEGORY_PROMPT = """A job seeker typed the job category they want: "{user_text}"
The text may be English, Marathi, Hindi or Hinglish, and may contain spelling mistakes.

Return the single best matching slug from this list, or null if nothing fits reasonably:
{category_list}

Examples: "computer operator" -> data-entry, "बँक" -> banking, "sarkari naukri" -> govt-jobs,
"nurse" -> healthcare, "fitter" -> iti-mechanical, "tally" -> accounts-finance.
Return only JSON: {{"slug": "..."}}"""

_JOB_TEXT_PROMPT = """Extract every distinct job opening from the text below for an Indian job-alert service for
students in Maharashtra. The text may be a WhatsApp forward, website copy or a list of jobs.

Rules:
- One object per distinct job. If the same post lists multiple roles, create one object per role.
- Do not invent company names, apply links, salaries or dates. Use null when missing.
- category_slug must be one of:
{category_list}
- job_type: "govt" for government / PSU / bank recruitment notifications, "internship" for internships
  and apprenticeships, "wfh" for fully remote, else "private".
- district: a Maharashtra district name in English if the job is in one district; null if remote,
  all-India, or multiple districts. Put the human-readable place in location_text.
- last_date: YYYY-MM-DD if a last date / closing date is written, else null.
- apply_link: only a real URL present in the text.
- description: at most 300 characters with the most useful details (vacancies, age limit, fee, timing).
- Ignore spam, "pay registration fee" offers and posts asking for money; do not output them.

TEXT:
{raw_text}

Return only JSON matching the schema."""


def _category_list_text(categories: dict[str, str]) -> str:
    return "\n".join(f"- {slug}: {name}" for slug, name in categories.items())


@lru_cache
def _get_client():
    from google import genai

    return genai.Client(api_key=settings.GEMINI_API_KEY)


async def _generate(contents: list, schema: type[BaseModel]) -> Optional[BaseModel]:
    if not settings.GEMINI_API_KEY:
        logger.warning("GEMINI_API_KEY not set; skipping Gemini call")
        return None
    try:
        from google.genai import types

        client = _get_client()
        response = await client.aio.models.generate_content(
            model=settings.GEMINI_MODEL,
            contents=contents,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=schema,
            ),
        )
    except Exception:  # noqa: BLE001
        logger.exception("Gemini generate_content call failed")
        return None

    parsed = getattr(response, "parsed", None)
    if isinstance(parsed, schema):
        return parsed
    try:
        data = json.loads(response.text)
        return schema.model_validate(data)
    except Exception:  # noqa: BLE001
        logger.exception("Failed to parse Gemini response as %s", schema.__name__)
        return None


def _read_docx_text(data: bytes) -> str:
    import io

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)[:MAX_JOB_TEXT_CHARS]


async def parse_resume(
    data: bytes, mime: str, categories: dict[str, str]
) -> Optional[ResumeData]:
    """Parse a resume file (PDF / image / DOCX) into structured data via Gemini."""
    prompt = _RESUME_PROMPT.format(category_list=_category_list_text(categories))

    if mime == "application/pdf" or mime.startswith("image/"):
        from google.genai import types

        contents = [prompt, types.Part.from_bytes(data=data, mime_type=mime)]
    # Legacy binary .doc (application/msword) is deliberately not accepted: python-docx only
    # reads the zipped-XML .docx format. The bot rejects .doc uploads before they get here.
    elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            text = await _to_thread(_read_docx_text, data)
        except Exception:  # noqa: BLE001
            logger.exception("Failed to extract text from DOCX")
            return None
        contents = [f"{prompt}\n\nRESUME TEXT:\n{text}"]
    else:
        logger.warning("Unsupported resume mime type for AI parsing: %s", mime)
        return None

    result = await _generate(contents, ResumeData)
    if result is None:
        return None
    result.suggested_category_slugs = [
        s for s in result.suggested_category_slugs if s in categories
    ][:5]
    result.skills = result.skills[:12]
    return result


async def map_category(user_text: str, categories: dict[str, str]) -> Optional[str]:
    prompt = _CATEGORY_PROMPT.format(
        user_text=user_text, category_list=_category_list_text(categories)
    )
    result = await _generate([prompt], CategoryMatch)
    if result is None or result.slug not in categories:
        return None
    return result.slug


async def parse_job_text(raw_text: str, categories: dict[str, str]) -> list[JobDraft]:
    prompt = _JOB_TEXT_PROMPT.format(
        category_list=_category_list_text(categories),
        raw_text=raw_text[:MAX_JOB_TEXT_CHARS],
    )
    result = await _generate([prompt], JobDraftList)
    if result is None:
        return []
    drafts: list[JobDraft] = []
    for draft in result.jobs:
        if draft.category_slug not in categories:
            draft.category_slug = "other"
        if draft.job_type not in ("govt", "private", "internship", "wfh"):
            draft.job_type = "private"
        drafts.append(draft)
    return drafts


async def _to_thread(func, *args):
    import asyncio

    return await asyncio.to_thread(func, *args)
